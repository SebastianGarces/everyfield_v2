"""Check native catalog exports after generation or preview download.

Run with the bundled artifact Python runtime (pypdf, openpyxl, python-docx):
  python3 scripts/check-document-exports.py /tmp/document-catalog/normal
This checks file structure and content; rendered-page inspection is still required.
"""
import sys
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

root = Path(sys.argv[1])
expected = {
    "commitment-card.pdf": 1,
    "response-card.pdf": 1,
    "guest-sign-in-sheet.pdf": 1,
    "vision-meeting-agenda.pdf": 1,
    "launch-sunday-checklists.pdf": 6,
}
for filename, pages in expected.items():
    pdf = PdfReader(root / filename)
    assert len(pdf.pages) == pages, filename
    size = (396, 306) if "card" in filename else (612, 792)
    for page in pdf.pages:
        assert (float(page.mediabox.width), float(page.mediabox.height)) == size
        assert page.extract_text().strip(), filename
    print(f"PASS {filename}: {pages} pages, correct page geometry, nonempty text")

for stem in ["vision-meeting-agenda", "board-meeting-agenda", "member-expectations", "launch-team-commitment", "follow-up-letter"]:
    document = Document(root / f"{stem}.docx")
    for section in document.sections:
        assert round(section.page_width.inches, 2) == 8.5
        assert round(section.page_height.inches, 2) == 11
        assert section.left_margin.inches >= 0.5
        assert section.right_margin.inches >= 0.5
    text = "\n".join(p.text for p in document.paragraphs)
    assert document.paragraphs[0].style.name == "Title"
    if "agenda" in stem:
        assert len(document.tables) == 1
        assert len(document.tables[0].rows) == (6 if stem == "vision-meeting-agenda" else 8)
    if stem in ["member-expectations", "launch-team-commitment"]:
        assert "Signature" in text and "Printed name" in text
    print(f"PASS {stem}.docx: Letter margins, title, required structure")

for stem, net_row, last_col in [("budget-worksheet", 18, "C"), ("first-year-budget", 19, "N")]:
    filename = root / f"{stem}.xlsx"
    sheet = load_workbook(filename).active
    assert sheet.freeze_panes == "B4"
    assert sheet[f"B{net_row}"].value.startswith("=B")
    assert "-B" in sheet[f"B{net_row}"].value
    assert sheet[f"{last_col}{net_row}"].data_type == "f"
    for row in sheet.iter_rows(min_row=4, min_col=2):
        for cell in row:
            if cell.data_type == "f":
                assert "$" in cell.number_format, cell.coordinate
    cached = load_workbook(filename, data_only=True).active
    assert cached[f"{last_col}{net_row}"].value == 0
    assert len(sheet.conditional_formatting) == 1
    # Regression: relative print-area rows make LibreOffice print rows outside
    # the selected area. Inspect the actual OOXML, not a reader's normalization.
    with ZipFile(filename) as archive:
        xml = ET.fromstring(archive.read("xl/workbook.xml"))
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    area = xml.find("s:definedNames/s:definedName[@name='_xlnm.Print_Area']", ns).text
    first_row = 3 if stem == "first-year-budget" else 1
    assert area.endswith(f"!$A${first_row}:${last_col}${net_row}"), area
    print(f"PASS {stem}.xlsx: frozen context, currency formulas, cached zeros, deficit rule, absolute print area")
